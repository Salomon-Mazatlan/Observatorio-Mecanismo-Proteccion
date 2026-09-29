#!/usr/bin/env python3
"""Update datos/comparativo_marco_legal.json from plantillas/comparativo_marco_legal.xlsx.

Usage (from the repository root):
    python herramientas/actualizar_comparativo.py plantillas/comparativo_marco_legal.xlsx
    python herramientas/actualizar_comparativo.py plantillas/comparativo_marco_legal.xlsx --solo-validar

Every sheet whose name is "federal" or a two-digit state code is read as an instrument.
The indicator list and value labels are kept from the current JSON.
"""
import argparse
import re

from openpyxl import load_workbook

from comun import RAIZ, catalogo_entidades, escribir_json, leer_json, respaldar, texto, verificar_raiz

RUTA = RAIZ / "datos/comparativo_marco_legal.json"
META = ["id", "cve_ent", "corto", "nombre", "publicacion", "ultima_reforma", "organo", "url"]
VALORES = {"si", "sí", "parcial", "no"}


def main():
    ap = argparse.ArgumentParser(description="Actualiza el comparativo legal desde Excel")
    ap.add_argument("excel")
    ap.add_argument("--solo-validar", action="store_true")
    args = ap.parse_args()
    verificar_raiz()

    actual = leer_json(RUTA)
    ids = {i["id"] for i in actual["indicadores"]} | {s["id"] for i in actual["indicadores"] for s in i["subindicadores"]}
    ents = dict(catalogo_entidades())
    wb = load_workbook(args.excel, data_only=True)
    errores, instrumentos, codificacion = [], [], {}

    for ws in wb.worksheets:
        # sheet names: federal, a state code, or a state code with a suffix (05-periodistas)
        if not (ws.title == "federal" or re.match(r"^\d{2}(-[a-z]+)?$", ws.title)):
            continue
        filas = list(ws.iter_rows(values_only=True))
        meta = {texto(f[0]): texto(f[1]) if len(f) > 1 else "" for f in filas[:len(META)] if f and f[0]}
        cve = "00" if ws.title == "federal" else ws.title[:2]
        if ws.title != "federal" and cve not in ents:
            errores.append(f"hoja {ws.title}: no es una clave de entidad válida")
            continue
        inst = {k: meta.get(k, "") for k in META}
        inst["id"] = ws.title
        inst["cve_ent"] = cve
        if not inst["corto"]:
            inst["corto"] = "Federal" if ws.title == "federal" else ents[cve]
        instrumentos.append(inst)
        cod = {}
        for n, f in enumerate(filas[len(META) + 2:], start=len(META) + 3):
            if not f or not f[0] or texto(f[0]) == "id":
                continue
            id_ = texto(f[0])
            if id_ not in ids:
                errores.append(f"hoja {ws.title}, fila {n}: indicador '{id_}' desconocido")
                continue
            valor = texto(f[2]).lower().replace("sí", "si") if len(f) > 2 else ""
            if valor and valor not in VALORES:
                errores.append(f"hoja {ws.title}, fila {n}: valor '{valor}' no permitido")
                continue
            if valor:
                cod[id_] = {"valor": valor, "articulos": texto(f[3]) if len(f) > 3 else "", "nota": texto(f[4]) if len(f) > 4 else ""}
        codificacion[ws.title] = cod

    definiciones = dict(actual.get("definiciones", {}))
    if "definiciones" in wb.sheetnames:
        for f in list(wb["definiciones"].iter_rows(values_only=True))[1:]:
            if f and f[0]:
                definiciones[texto(f[0])] = {"periodista": texto(f[1]) if len(f) > 1 else "", "defensora": texto(f[2]) if len(f) > 2 else "", "agresion": texto(f[3]) if len(f) > 3 else ""}

    if errores:
        print("Errores encontrados, no se escribió nada:")
        print("\n".join("  " + e for e in errores))
        raise SystemExit(1)
    print(f"{len(instrumentos)} instrumentos; celdas codificadas: " + ", ".join(f"{k}={len(v)}" for k, v in codificacion.items()))
    if args.solo_validar:
        print("Modo solo validar, no se escribió nada.")
        return
    resp = respaldar(RUTA)
    if resp:
        print(f"Respaldo: {resp}")
    actual["instrumentos"] = instrumentos
    actual["codificacion"] = codificacion
    actual["definiciones"] = definiciones
    escribir_json(RUTA, actual)
    print(f"Escrito {RUTA}")


if __name__ == "__main__":
    main()
