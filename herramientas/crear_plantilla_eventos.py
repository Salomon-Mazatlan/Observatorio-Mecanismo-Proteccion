#!/usr/bin/env python3
"""Build plantillas/eventos.xlsx from datos/eventos.json and datos/fuentes.json.

Usage (from the repository root):
    python herramientas/crear_plantilla_eventos.py [salida.xlsx]
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import Rule
from openpyxl.styles.differential import DifferentialStyle

from comun import DIR_PLANTILLAS, RUTA_EVENTOS, RUTA_FUENTES, verificar_raiz

COLS = ["id", "pais", "persona", "tema", "grupo", "subtema", "genero", "labor", "tipo", "fecha", "lat", "lon", "lugar", "titulo", "descripcion", "fuente", "url", "verificacion", "ejemplo"]
ANCHOS = [9, 7, 30, 16, 12, 16, 14, 14, 22, 12, 10, 11, 30, 48, 70, 16, 40, 15, 9]
GENEROS = ["femenino", "masculino", "lgbt", "no determinado", "no aplica"]
LABORES = ["busqueda", "ambiental", "indigena", "mujeres", "lgbt", "migracion", "civil", "animales", "organizacion"]
GRUPOS = ["periodista", "defensor", "ambos"]
SUBTEMAS = ["asesinato", "desaparicion", "agresion", "amenaza", "acoso_judicial", "ataque_medio"]
TEMAS = ["migracion", "desplazamiento", "desaparicion", "periodistas", "contexto"]
VERIF = ["oficial", "organización", "campo", "prensa", "sin verificar"]
MAX_FILAS = 1000

INSTRUCCIONES = [
    "Plantilla de captura de eventos del observatorio",
    "",
    "Captura en la hoja 'eventos' (celdas amarillas). Cada fila es un marcador en el mapa.",
    "Las filas con ejemplo = si son plantillas; bórralas o cambia ejemplo = no cuando las sustituyas por hechos reales.",
    "",
    "id: déjalo vacío en registros nuevos, el script lo asigna (ev-001, ev-002...). Para corregir un evento ya publicado conserva su id.",
    "tema: migracion, desplazamiento, desaparicion, periodistas o contexto (lista desplegable). Los tres primeros están desactivados en el sitio pero se conservan.",
    "grupo: solo para el tema periodistas; periodista, defensor o ambos (por ejemplo, un ataque a un medio que también cubre a defensores).",
    "pais: MX (México, predeterminado si se deja vacío), HN (Honduras) o CO (Colombia); las coordenadas se validan contra el país.",
    "genero: femenino, masculino, lgbt, no determinado (persona no identificada) o no aplica (comunidades, organizaciones o medios).",
    "labor: solo para personas defensoras; busqueda, ambiental, indigena, mujeres, lgbt, migracion, civil, animales u organizacion. Alimenta los chips de 'Labor de la persona defensora'.",
    "subtema: tipo de agresión o violencia; asesinato, desaparicion, agresion, amenaza, acoso_judicial o ataque_medio. Los informes y cifras agregadas no son eventos: van en datos/informes.json. Alimenta los chips de 'Tipo de agresión'.",
    "tipo: texto libre pero consistente (asesinato de periodista, asesinato de persona defensora, desaparición, agresión física, amenaza, acoso judicial, ataque a domicilio u oficina...).",
    "fecha: formato AAAA-MM-DD. Si solo se conoce el mes, usa el día 01 y anótalo en descripcion.",
    "lat, lon: grados decimales. En Google Maps clic derecho sobre el punto y copiar coordenadas. Si solo se conoce el municipio, usa la cabecera y anota 'ubicación aproximada' en descripcion.",
    "lugar: localidad, municipio y estado, por ejemplo 'Tepuche, Culiacán, Sinaloa'. El nombre del estado se usa para contar eventos por entidad.",
    "persona: nombre completo de la persona afectada, escrito siempre igual en todos sus eventos (así se agrupan y se detectan duplicados; la columna resalta en rojo los nombres repetidos). Varias personas se separan con punto y coma. Se deja vacío en colectivos sin nombre y en agresiones a personas vivas que no se nombran por criterio de cuidado.",
    "titulo: una línea. descripcion: qué pasó según la fuente, con cifras y quién las dio, sin datos que identifiquen a personas.",
    "fuente: id del catálogo (hoja 'catalogos'). Para agregar una fuente nueva hay que darla de alta en datos/fuentes.json y regenerar esta plantilla.",
    "url: enlace a la nota o documento, si existe. Cada fila es un caso individual; los informes y cifras agregadas van en datos/informes.json.",
    "verificacion: oficial (la cifra la dio una autoridad), organización (la documentó una organización civil), campo, prensa (solo la nota) o sin verificar.",
    "ejemplo: si / no.",
    "",
    "Para publicar: guarda este archivo y ejecuta desde la raíz del repositorio",
    "    python herramientas/actualizar_eventos.py plantillas/eventos.xlsx",
    "El script valida las filas, actualiza datos/eventos.json y deja un respaldo del JSON anterior. Luego sube datos/eventos.json a GitHub.",
    "Para regenerar esta plantilla con los eventos y fuentes actuales: python herramientas/crear_plantilla_eventos.py",
]


def main():
    verificar_raiz()
    DIR_PLANTILLAS.mkdir(exist_ok=True)
    salida = Path(sys.argv[1]) if len(sys.argv) > 1 else DIR_PLANTILLAS / "eventos.xlsx"
    eventos = json.loads(RUTA_EVENTOS.read_text(encoding="utf-8"))
    fuentes = json.loads(RUTA_FUENTES.read_text(encoding="utf-8"))
    arial, bold = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
    amarillo, azul = PatternFill("solid", fgColor="FFF2CC"), PatternFill("solid", fgColor="D9E2F3")

    wb = Workbook()
    we = wb.active
    we.title = "eventos"
    we.append(COLS)
    for e in sorted(eventos, key=lambda x: (x["fecha"], x["id"])):
        we.append([("si" if e[c] else "no") if c == "ejemplo" else e.get(c, "") for c in COLS])
    for i, w in enumerate(ANCHOS, 1):
        we.column_dimensions[get_column_letter(i)].width = w
    for row in we.iter_rows():
        for c in row:
            c.font = arial
            c.alignment = Alignment(vertical="top", wrap_text=c.column in (3, 13, 14, 15))
    for c in we[1]:
        c.font, c.fill = bold, azul
    for r in range(2, MAX_FILAS + 2):
        for ci in range(1, len(COLS) + 1):
            we.cell(r, ci).fill = amarillo
        we.cell(r, 4).number_format = "yyyy-mm-dd"
    we.freeze_panes = "B2"

    wc = wb.create_sheet("catalogos")
    wc.append(["tema", "verificacion", "fuente_id", "fuente_nombre", "fuente_tipo", "grupo", "subtema", "genero", "labor"])
    n = max(len(TEMAS), len(VERIF), len(fuentes), len(GRUPOS), len(SUBTEMAS), len(GENEROS), len(LABORES))
    for i in range(n):
        f = fuentes[i] if i < len(fuentes) else None
        wc.append([TEMAS[i] if i < len(TEMAS) else None, VERIF[i] if i < len(VERIF) else None,
                   f["id"] if f else None, f["nombre"] if f else None, f["tipo"] if f else None, GRUPOS[i] if i < len(GRUPOS) else None,
                   SUBTEMAS[i] if i < len(SUBTEMAS) else None, GENEROS[i] if i < len(GENEROS) else None, LABORES[i] if i < len(LABORES) else None])
    for w, col in zip([16, 16, 18, 80, 14, 12, 16, 16, 14], "ABCDEFGHI"):
        wc.column_dimensions[col].width = w
    for row in wc.iter_rows():
        for c in row:
            c.font = arial
    for c in wc[1]:
        c.font, c.fill = bold, azul

    rango = f"2:{MAX_FILAS + 1}"
    reglas = [
        DataValidation(type="list", formula1=f"=catalogos!$A$2:$A${len(TEMAS) + 1}", allow_blank=True), "D",
        DataValidation(type="list", formula1=f"=catalogos!$F$2:$F${len(GRUPOS) + 1}", allow_blank=True), "E",
        DataValidation(type="list", formula1=f"=catalogos!$G$2:$G${len(SUBTEMAS) + 1}", allow_blank=True), "F",
        DataValidation(type="list", formula1=f"=catalogos!$H$2:$H${len(GENEROS) + 1}", allow_blank=True), "G",
        DataValidation(type="list", formula1=f"=catalogos!$I$2:$I${len(LABORES) + 1}", allow_blank=True), "H",
        DataValidation(type="list", formula1=f"=catalogos!$B$2:$B${len(VERIF) + 1}", allow_blank=True), "R",
        DataValidation(type="list", formula1=f"=catalogos!$C$2:$C${len(fuentes) + 1}", allow_blank=True), "P",
        DataValidation(type="list", formula1='"si,no"', allow_blank=True), "S",
        DataValidation(type="decimal", operator="between", formula1="-4.3", formula2="33", allow_blank=True), "K",
        DataValidation(type="decimal", operator="between", formula1="-119", formula2="-66.8", allow_blank=True), "L",
    ]
    for dv, col in zip(reglas[::2], reglas[1::2]):
        dv.error, dv.showErrorMessage = "Valor no permitido", True
        dv.add(f"{col}{rango.replace(':', ':' + col)}")
        we.add_data_validation(dv)
    # Repeated names in the persona column turn red, to spot duplicates or several events of one person
    rojo = DifferentialStyle(font=Font(color="9C0006"), fill=PatternFill("solid", start_color="FFC7CE"))
    we.conditional_formatting.add(f"C2:C{MAX_FILAS + 1}", Rule(type="duplicateValues", dxf=rojo))

    wi = wb.create_sheet("instrucciones")
    for t in INSTRUCCIONES:
        wi.append([t])
    for row in wi.iter_rows():
        for c in row:
            c.font = bold if c.row == 1 else arial
    wi.column_dimensions["A"].width = 130
    wb.save(salida)
    print(f"Plantilla escrita en {salida} con {len(eventos)} eventos y {len(fuentes)} fuentes")


if __name__ == "__main__":
    main()
