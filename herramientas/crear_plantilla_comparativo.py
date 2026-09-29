#!/usr/bin/env python3
"""Build plantillas/comparativo_marco_legal.xlsx from datos/comparativo_marco_legal.json.

Usage (from the repository root):
    python herramientas/crear_plantilla_comparativo.py

One sheet per instrument (federal, one per state coded), with a row per indicator
or subindicator: valor (si / parcial / no), articulos, nota. Add a state by copying
a sheet, renaming it to the two-digit INEGI code and filling the header rows.
"""
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation

from comun import AMARILLO, AZUL, BOLD, DIR_PLANTILLAS, RAIZ, estilizar_hoja, hoja_instrucciones, leer_json, verificar_raiz

RUTA = RAIZ / "datos/comparativo_marco_legal.json"
INSTRUCCIONES = [
    "Plantilla del comparativo de instrumentos de protección a periodistas y personas defensoras",
    "",
    "Hay una hoja por instrumento (federal y un código de dos dígitos por estado). Las primeras filas son los metadatos del instrumento; después va una fila por indicador o subindicador.",
    "valor: si, parcial o no (lista desplegable). articulos: dónde se ubica en el texto (por ejemplo 'Arts. 2, 24'). nota: qué dice y qué le falta, en una o dos frases.",
    "Los subindicadores están sangrados y sirven para desglosar definiciones y medidas; no cambian el valor del indicador principal, que se codifica por separado.",
    "Para agregar un estado: copia la hoja de un estado, renómbrala con su clave INEGI (por ejemplo 14 para Jalisco) y llena los metadatos y las celdas. Si el estado tiene más de una ley, usa la clave con un sufijo (05-periodistas, 05-defensoras). Las hojas 'hn' y 'co' son las normas nacionales de Honduras y Colombia.",
    "La hoja 'definiciones' guarda el texto literal de las definiciones de periodista, persona defensora y agresión de cada instrumento, que alimenta el análisis de definiciones.",
    "",
    "Para publicar: python herramientas/actualizar_comparativo.py plantillas/comparativo_marco_legal.xlsx",
    "El script valida los valores, reescribe datos/comparativo_marco_legal.json completo y deja respaldo.",
]
META = ["id", "cve_ent", "pais", "corto", "nombre", "publicacion", "ultima_reforma", "organo", "url"]


def main():
    verificar_raiz()
    cmp = leer_json(RUTA)
    filas_ind = []
    for ind in cmp["indicadores"]:
        filas_ind.append((ind["id"], f"{ind['numero']}. {ind['nombre']}", False))
        for s in ind["subindicadores"]:
            filas_ind.append((s["id"], s["nombre"], True))

    wb = Workbook()
    wb.remove(wb.active)
    for inst in cmp["instrumentos"]:
        ws = wb.create_sheet(inst["id"])
        for k in META:
            ws.append([k, inst.get(k, "")])
        ws.append([])
        ws.append(["id", "indicador", "valor", "articulos", "nota"])
        cod = cmp["codificacion"].get(inst["id"], {})
        for id_, nombre, sub in filas_ind:
            c = cod.get(id_, {})
            ws.append([id_, ("    " + nombre) if sub else nombre, c.get("valor", ""), c.get("articulos", ""), c.get("nota", "")])
        estilizar_hoja(ws, [8, 70, 10, 34, 90], filas_editables=0, columnas_ajustables=(2, 4, 5))
        ini = len(META) + 2
        for r in range(2, len(META) + 1):
            ws.cell(r, 2).fill = AMARILLO
        for r in range(ini + 1, ini + 1 + len(filas_ind)):
            for ci in (3, 4, 5):
                ws.cell(r, ci).fill = AMARILLO
        for c in ws[ini]:
            c.font, c.fill = BOLD, AZUL
        dv = DataValidation(type="list", formula1='"si,parcial,no"', allow_blank=True)
        dv.error, dv.showErrorMessage = "Valor no permitido", True
        dv.add(f"C{ini + 1}:C{ini + len(filas_ind)}")
        ws.add_data_validation(dv)
        ws.freeze_panes = f"C{ini + 1}"

    wd = wb.create_sheet("definiciones")
    wd.append(["instrumento", "periodista", "defensora", "agresion"])
    for inst in cmp["instrumentos"]:
        d = cmp["definiciones"].get(inst["id"], {})
        wd.append([inst["id"], d.get("periodista", ""), d.get("defensora", ""), d.get("agresion", "")])
    estilizar_hoja(wd, [12, 70, 70, 70], filas_editables=len(cmp["instrumentos"]), columnas_ajustables=(2, 3, 4))

    wi = wb.create_sheet("indicadores")
    wi.append(["id", "nombre", "es_subindicador"])
    for id_, nombre, sub in filas_ind:
        wi.append([id_, nombre.strip(), "si" if sub else "no"])
    estilizar_hoja(wi, [8, 100, 16])

    hoja_instrucciones(wb, INSTRUCCIONES)
    DIR_PLANTILLAS.mkdir(exist_ok=True)
    salida = DIR_PLANTILLAS / "comparativo_marco_legal.xlsx"
    wb.save(salida)
    print(f"{salida}: {len(cmp['instrumentos'])} instrumentos, {len(filas_ind)} filas por instrumento")


if __name__ == "__main__":
    main()
