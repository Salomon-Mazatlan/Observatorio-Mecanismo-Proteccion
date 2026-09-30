#!/usr/bin/env python3
"""Update datos/eventos.json from the Excel capture template.

Usage (run from the repository root):
    python herramientas/actualizar_eventos.py plantillas/eventos.xlsx
    python herramientas/actualizar_eventos.py plantillas/eventos.xlsx --reemplazar
    python herramientas/actualizar_eventos.py plantillas/eventos.xlsx --solo-validar

Default mode merges by id: existing ids are updated, new ids are added,
ids missing from the Excel file are kept. --reemplazar rewrites the whole
file with the Excel content. A backup of the previous JSON is always kept.
"""
import argparse
import json
import re
import shutil
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

RUTA_EVENTOS = Path("datos/eventos.json")
RUTA_FUENTES = Path("datos/fuentes.json")
HOJA = "eventos"

TEMAS = {"migracion", "desplazamiento", "desaparicion", "periodistas", "contexto"}
GRUPOS = {"periodista", "defensor", "ambos"}
GENEROS = {"femenino", "masculino", "lgbt", "no determinado", "no aplica"}
LABORES = {"busqueda", "ambiental", "indigena", "mujeres", "lgbt", "migracion", "civil", "animales", "organizacion"}
SUBTEMAS = {"asesinato", "desaparicion", "agresion", "amenaza", "acoso_judicial", "ataque_medio", "proteccion", "informe"}
VERIFICACION = {"oficial", "organización", "campo", "prensa", "sin verificar"}
# Bounding boxes per country (lat_min, lat_max, lon_min, lon_max)
LIMITES = {"MX": (14, 33, -119, -86), "HN": (12.9, 17.5, -89.4, -83.1), "CO": (-4.3, 13.5, -81.8, -66.8)}
COLUMNAS = ["id", "pais", "persona", "tema", "grupo", "subtema", "genero", "labor", "tipo", "fecha", "lat", "lon", "lugar", "titulo",
            "descripcion", "fuente", "url", "verificacion", "ejemplo"]
OBLIGATORIAS = ["tema", "tipo", "fecha", "lat", "lon", "lugar", "titulo", "fuente", "verificacion"]


def leer_excel(ruta):
    wb = load_workbook(ruta, data_only=True)
    if HOJA not in wb.sheetnames:
        sys.exit(f"La hoja '{HOJA}' no existe en {ruta}")
    ws = wb[HOJA]
    filas = list(ws.iter_rows(values_only=True))
    encabezados = [str(c).strip().lower() if c else "" for c in filas[0]]
    faltan = [c for c in COLUMNAS if c not in encabezados]
    if faltan:
        sys.exit(f"Faltan columnas en la hoja: {', '.join(faltan)}")
    idx = {c: encabezados.index(c) for c in COLUMNAS}
    registros = []
    for n, fila in enumerate(filas[1:], start=2):
        if all(c in (None, "") for c in fila):
            continue
        registros.append((n, {c: fila[idx[c]] for c in COLUMNAS}))
    return registros


def texto(v):
    return "" if v is None else str(v).strip()


def a_fecha(v):
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    s = texto(v)
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    return None


def a_bool(v):
    return texto(v).lower() in ("si", "sí", "true", "1", "x", "verdadero")


def validar(registros, fuentes_ids):
    errores, limpios, ids = [], [], set()
    for n, r in registros:
        e = []
        for c in OBLIGATORIAS:
            if texto(r[c]) == "":
                e.append(f"falta '{c}'")
        tema = texto(r["tema"]).lower()
        if tema and tema not in TEMAS:
            e.append(f"tema '{tema}' no válido")
        grupo = texto(r["grupo"]).lower()
        if tema == "periodistas" and grupo not in GRUPOS:
            e.append("grupo debe ser periodista, defensor o ambos para el tema periodistas")
        subtema = texto(r["subtema"]).lower()
        if subtema and subtema not in SUBTEMAS:
            e.append(f"subtema '{subtema}' no válido")
        genero = texto(r["genero"]).lower()
        if genero and genero not in GENEROS:
            e.append(f"genero '{genero}' no válido")
        labor = texto(r["labor"]).lower()
        if labor and labor not in LABORES:
            e.append(f"labor '{labor}' no válida")
        ver = texto(r["verificacion"]).lower()
        if ver and ver not in VERIFICACION:
            e.append(f"verificacion '{ver}' no válida")
        fuente = texto(r["fuente"])
        if fuente and fuente not in fuentes_ids:
            e.append(f"fuente '{fuente}' no está en fuentes.json")
        fecha = a_fecha(r["fecha"]) if texto(r["fecha"]) else None
        if texto(r["fecha"]) and not fecha:
            e.append("fecha no reconocida (usar AAAA-MM-DD)")
        pais = (texto(r["pais"]) or "MX").upper()
        if pais not in LIMITES:
            e.append(f"pais '{pais}' no válido (MX, HN o CO)")
            pais = "MX"
        try:
            lat, lon = float(r["lat"]), float(r["lon"])
            la0, la1, lo0, lo1 = LIMITES[pais]
            if not (la0 <= lat <= la1 and lo0 <= lon <= lo1):
                e.append(f"coordenadas fuera de {dict(MX='México', HN='Honduras', CO='Colombia')[pais]}")
        except (TypeError, ValueError):
            lat = lon = None
            e.append("lat/lon deben ser numéricos")
        rid = texto(r["id"])
        if rid and rid in ids:
            e.append(f"id '{rid}' repetido")
        ids.add(rid)
        if e:
            errores.append(f"Fila {n}: " + "; ".join(e))
            continue
        reg = {
            "id": rid, "tema": tema, "tipo": texto(r["tipo"]), "fecha": fecha,
            "lat": round(lat, 5), "lon": round(lon, 5), "lugar": texto(r["lugar"]),
            "titulo": texto(r["titulo"]), "descripcion": texto(r["descripcion"]),
            "fuente": fuente, "url": texto(r["url"]), "verificacion": ver,
            "ejemplo": a_bool(r["ejemplo"]),
        }
        if pais != "MX":
            reg["pais"] = pais
        # Names separated by semicolons, extra spaces removed
        reg["persona"] = "; ".join(" ".join(x.split()) for x in texto(r["persona"]).split(";") if x.strip())
        if grupo:
            reg["grupo"] = grupo
        if subtema:
            reg["subtema"] = subtema
        if genero:
            reg["genero"] = genero
        if labor:
            reg["labor"] = labor
        limpios.append(reg)
    return limpios, errores


def normalizar(nombre):
    import unicodedata
    s = unicodedata.normalize("NFD", nombre.lower())
    return " ".join("".join(c for c in s if unicodedata.category(c) != "Mn").split())


def avisos_personas(eventos):
    """Warn about names that repeat: likely duplicates or several events of the same person."""
    por_nombre = {}
    for e in eventos:
        for nombre in [x.strip() for x in e.get("persona", "").split(";") if x.strip()]:
            por_nombre.setdefault(normalizar(nombre), (nombre, []))[1].append(e)
    avisos = []
    for nombre, lista in por_nombre.values():
        if len(lista) < 2:
            continue
        mismos = {}
        for e in lista:
            mismos.setdefault(e.get("subtema", ""), []).append(e)
        dup = [v for k, v in mismos.items() if len(v) > 1 and k in ("asesinato", "desaparicion")]
        tipo = "POSIBLE DUPLICADO" if dup else "varios eventos"
        detalle = ", ".join(f"{e['id']} ({e['fecha']}, {e.get('subtema', '')})" for e in sorted(lista, key=lambda x: x["fecha"]))
        avisos.append(f"  {tipo}: {nombre} -> {detalle}")
    return avisos


def asignar_ids(nuevos, existentes):
    usados = {e["id"] for e in existentes} | {e["id"] for e in nuevos if e["id"]}
    numeros = [int(m.group(1)) for i in usados if (m := re.match(r"ev-(\d+)$", i))]
    siguiente = max(numeros, default=0) + 1
    for e in nuevos:
        if not e["id"]:
            e["id"] = f"ev-{siguiente:03d}"
            siguiente += 1
    return nuevos


def main():
    ap = argparse.ArgumentParser(description="Actualiza datos/eventos.json desde Excel")
    ap.add_argument("excel", help="archivo .xlsx de captura")
    ap.add_argument("--reemplazar", action="store_true", help="sustituir todo el JSON por el Excel")
    ap.add_argument("--solo-validar", action="store_true", help="revisar sin escribir")
    args = ap.parse_args()

    if not RUTA_FUENTES.exists():
        sys.exit("Ejecuta el script desde la raíz del repositorio (no encuentro datos/fuentes.json)")
    fuentes_ids = {f["id"] for f in json.loads(RUTA_FUENTES.read_text(encoding="utf-8"))}
    existentes = json.loads(RUTA_EVENTOS.read_text(encoding="utf-8")) if RUTA_EVENTOS.exists() else []

    nuevos, errores = validar(leer_excel(args.excel), fuentes_ids)
    if errores:
        print("Errores encontrados, no se escribió nada:")
        print("\n".join("  " + e for e in errores))
        sys.exit(1)
    nuevos = asignar_ids(nuevos, existentes)

    if args.reemplazar:
        resultado = nuevos
        agregados, actualizados = len(nuevos), 0
    else:
        por_id = {e["id"]: e for e in existentes}
        agregados = sum(1 for e in nuevos if e["id"] not in por_id)
        actualizados = len(nuevos) - agregados
        por_id.update({e["id"]: e for e in nuevos})
        resultado = list(por_id.values())
    resultado.sort(key=lambda e: (e["fecha"], e["id"]))

    print(f"Filas válidas: {len(nuevos)}. Agregados: {agregados}. Actualizados: {actualizados}. Total: {len(resultado)}.")
    avisos = avisos_personas(resultado)
    if avisos:
        print(f"Nombres repetidos ({len(avisos)}); revisa si son la misma persona:")
        print("\n".join(avisos))
    if args.solo_validar:
        print("Modo solo validar, no se escribió nada.")
        return
    if RUTA_EVENTOS.exists():
        respaldo = RUTA_EVENTOS.with_name(f"eventos_{datetime.now():%Y%m%d_%H%M%S}.bak.json")
        shutil.copy(RUTA_EVENTOS, respaldo)
        print(f"Respaldo: {respaldo}")
    RUTA_EVENTOS.write_text(json.dumps(resultado, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Escrito {RUTA_EVENTOS}")


if __name__ == "__main__":
    main()
